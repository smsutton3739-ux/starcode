"""Exporting a report to PDF, DOCX, CSV, Markdown or JSON.

Every format carries the claim-type labelling. An exported PDF that dropped the
distinction between a calculation and a conjecture would defeat the entire product, so
the label travels with the claim into every format — including the flat CSV, where it is
a column.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone
from typing import Any

from app.agents.contracts import CLAIM_TYPE_PRESENTATION
from app.db.models.analysis import Analysis, ClaimType

EXPORT_FORMATS = ["pdf", "docx", "csv", "markdown", "json"]

MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "csv": "text/csv",
    "markdown": "text/markdown",
    "json": "application/json",
}

DISCLAIMER = (
    "Every statement in this report is labelled by kind. Statements marked as an AI "
    "hypothesis, a traditional interpretation, or a scholarly interpretation are not "
    "established fact and must not be cited as such. Astronomical and calendrical "
    "calculations are reproducible but carry the stated accuracy of the algorithm that "
    "produced them. This report is a research aid, not a scholarly authority."
)

#: Office core properties cap each field at 255 characters. The full disclaimer always
#: appears in the document body; this is only the metadata summary.
SHORT_DISCLAIMER = (
    "Claims are labelled by kind. AI hypotheses and traditional or scholarly "
    "interpretations are not established fact. A research aid, not an authority."
)


class ExportError(RuntimeError):
    pass


def _report_of(analysis: Analysis) -> Any:
    if not analysis.reports:
        raise ExportError(
            "This analysis has no report yet. Wait for it to finish before exporting."
        )
    return max(analysis.reports, key=lambda r: r.version)


def _label(claim_type: ClaimType) -> str:
    return CLAIM_TYPE_PRESENTATION[claim_type]["label"]


def _sorted_claims(analysis: Analysis) -> list:
    return sorted(analysis.claims, key=lambda c: c.ordering)


def export_analysis(analysis: Analysis, fmt: str) -> tuple[bytes, str, str]:
    """Returns (content, mime_type, filename)."""
    if fmt not in EXPORT_FORMATS:
        raise ExportError(f"Unsupported format {fmt!r}. Choose one of: {', '.join(EXPORT_FORMATS)}.")

    safe_title = "".join(
        c if c.isalnum() or c in " -_" else "_" for c in (analysis.title or "analysis")
    ).strip()[:60] or "analysis"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    filename = f"{safe_title}-{stamp}.{'md' if fmt == 'markdown' else fmt}"

    builder = {
        "markdown": _build_markdown,
        "json": _build_json,
        "csv": _build_csv,
        "pdf": _build_pdf,
        "docx": _build_docx,
    }[fmt]

    return builder(analysis), MIME_TYPES[fmt], filename


# --------------------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------------------


def _build_markdown(analysis: Analysis) -> bytes:
    report = _report_of(analysis)
    out: list[str] = []

    out.append(f"# {analysis.title}\n")
    out.append(f"*Generated {datetime.now(timezone.utc):%d %B %Y} by Starcode*\n")
    out.append(f"> {DISCLAIMER}\n")

    confidence = report.confidence_summary or {}
    if confidence.get("overall") is not None:
        out.append(
            f"**Overall confidence:** {confidence['overall']:.0%} "
            f"({confidence.get('band', 'unrated')})\n"
        )
        out.append(f"{confidence.get('rationale', '')}\n")

    out.append("## Executive Summary\n")
    out.append(report.executive_summary + "\n")

    for section in report.sections:
        if section["key"] == "executive_summary":
            continue
        out.append(f"## {section['title']}\n")

        if section.get("is_empty"):
            out.append(f"*{section.get('empty_reason', 'No findings.')}*\n")
            continue

        for claim in section.get("claims", []):
            label = claim.get("presentation", {}).get("label", claim["claim_type"])
            out.append(f"**[{label}]** {claim['statement']}")
            out.append("")
            if claim.get("quoted_text") and claim["claim_type"] == "source_text":
                quoted = claim["quoted_text"].replace("\n", "\n> ")
                out.append(f"> {quoted}\n")
            if claim.get("reasoning"):
                out.append(f"*Reasoning:* {claim['reasoning']}\n")
            out.append(
                f"*Confidence:* {claim['confidence']:.0%}"
                + (f" — {claim['confidence_basis']}" if claim.get("confidence_basis") else "")
                + "\n"
            )
            if claim.get("engine"):
                out.append(f"*Computed by:* {claim['engine']}")
                if claim.get("algorithm_reference"):
                    out.append(f" ({claim['algorithm_reference']})")
                out.append("\n")
            for reference in claim.get("references", []):
                mark = "" if reference.get("verified") else " *(unverified citation)*"
                out.append(f"- {reference['citation_text']}{mark}\n")
            out.append("")

        data = section.get("data") or {}
        if section["key"] == "timeline" and data.get("events"):
            out.append("| Year | Event | Certainty |")
            out.append("|---|---|---|")
            for event in data["events"]:
                out.append(
                    f"| {event.get('label', event.get('year'))} | "
                    f"{event.get('description', '')} | {event.get('certainty', '')} |"
                )
            out.append("")
        if section["key"] == "references" and data.get("references"):
            for reference in data["references"]:
                mark = "" if reference.get("verified") else " *(unverified — check before citing)*"
                out.append(f"- {reference['citation_text']}{mark}")
            out.append("")
        if section["key"] == "further_reading" and data.get("items"):
            for item in data["items"]:
                out.append(f"- **{item['title']}** — {item['citation']}  \n  {item['why']}")
            out.append("")

    out.append("## How this report was produced\n")
    for step in report.reasoning_trace or []:
        out.append(
            f"{step['step']}. **{step['agent']}** ({step['status']}"
            + (f", {step['model']}" if step.get("model") else "")
            + f") — {step.get('reasoning', '')}"
        )
    out.append("")

    return "\n".join(out).encode("utf-8")


# --------------------------------------------------------------------------------------
# JSON
# --------------------------------------------------------------------------------------


def _build_json(analysis: Analysis) -> bytes:
    report = _report_of(analysis)
    payload = {
        "analysis": {
            "id": analysis.id,
            "title": analysis.title,
            "status": analysis.status.value,
            "created_at": analysis.created_at.isoformat(),
            "completed_at": analysis.completed_at.isoformat() if analysis.completed_at else None,
            "overall_confidence": analysis.overall_confidence,
            "confidence_rationale": analysis.confidence_rationale,
            "detected_language": analysis.detected_language,
            "source_identification": analysis.source_identification,
            "failed_stages": list(analysis.failed_stages or []),
        },
        "disclaimer": DISCLAIMER,
        "claim_type_vocabulary": {
            key.value: value for key, value in CLAIM_TYPE_PRESENTATION.items()
        },
        "report": {
            "executive_summary": report.executive_summary,
            "sections": report.sections,
            "timeline": report.timeline,
            "confidence_summary": report.confidence_summary,
            "further_reading": report.further_reading,
            "reasoning_trace": report.reasoning_trace,
            "generator_version": report.generator_version,
        },
        "claims": [
            {
                "id": c.id,
                "section": c.section,
                "claim_type": c.claim_type.value,
                "claim_type_label": _label(c.claim_type),
                "statement": c.statement,
                "reasoning": c.reasoning,
                "confidence": c.confidence,
                "confidence_basis": c.confidence_basis,
                "produced_by": c.produced_by,
                "engine": c.engine,
                "algorithm_reference": c.algorithm_reference,
                "quoted_text": c.quoted_text,
                "payload": c.payload,
            }
            for c in _sorted_claims(analysis)
        ],
        "entities": [
            {
                "name": e.name,
                "entity_type": e.entity_type.value,
                "description": e.description,
                "extraction_confidence": e.extraction_confidence,
                "identification_confidence": e.identification_confidence,
                "mention_count": e.mention_count,
                "earliest_year": e.earliest_year,
                "latest_year": e.latest_year,
            }
            for e in analysis.entities
        ],
        "references": [
            {
                "citation_text": r.citation_text,
                "url": r.url,
                "reference_kind": r.reference_kind,
                "verified": r.verified,
            }
            for r in analysis.references
        ],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False, default=str).encode("utf-8")


# --------------------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------------------


def _build_csv(analysis: Analysis) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, quoting=csv.QUOTE_ALL)
    writer.writerow(
        [
            "section", "claim_type", "claim_type_label", "is_evidence", "statement",
            "reasoning", "confidence", "confidence_basis", "produced_by", "engine",
            "algorithm_reference", "quoted_text", "citations",
        ]
    )

    evidence_types = {"source_text", "verified_history", "astronomical_calculation"}
    references_by_claim: dict[str, list[str]] = {}
    for reference in analysis.references:
        if reference.claim_id:
            references_by_claim.setdefault(reference.claim_id, []).append(
                reference.citation_text + ("" if reference.verified else " [unverified]")
            )

    for claim in _sorted_claims(analysis):
        writer.writerow(
            [
                claim.section,
                claim.claim_type.value,
                _label(claim.claim_type),
                "yes" if claim.claim_type.value in evidence_types else "no",
                claim.statement,
                claim.reasoning or "",
                f"{claim.confidence:.3f}",
                claim.confidence_basis or "",
                claim.produced_by,
                claim.engine or "",
                claim.algorithm_reference or "",
                (claim.quoted_text or "")[:2000],
                " | ".join(references_by_claim.get(claim.id, [])),
            ]
        )

    # Excel opens UTF-8 CSV as mojibake without a BOM, and this corpus is full of
    # transliterated names that would be mangled.
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


# --------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------


def _build_pdf(analysis: Analysis) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_JUSTIFY
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            HRFlowable,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover
        raise ExportError("PDF export requires reportlab, which is not installed.") from exc

    from xml.sax.saxutils import escape

    report = _report_of(analysis)
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        title=analysis.title,
        author="Starcode",
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.5, leading=13,
                          alignment=TA_JUSTIFY, spaceAfter=4)
    small = ParagraphStyle("Small", parent=body, fontSize=8, textColor=colors.HexColor("#555555"))
    quote = ParagraphStyle("Quote", parent=body, leftIndent=10 * mm, fontName="Times-Italic",
                           textColor=colors.HexColor("#333333"))
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=18, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, spaceBefore=10, spaceAfter=4)

    # Each claim type gets a colour, matching the web UI so a printed report reads the
    # same way as the screen.
    tone_colors = {
        "source_text": "#475569", "verified_history": "#15803d",
        "astronomical_calculation": "#1d4ed8", "textual_analysis": "#0f766e",
        "traditional_interpretation": "#a16207", "scholarly_interpretation": "#7c3aed",
        "ai_hypothesis": "#c2410c", "uncertain": "#64748b",
    }

    story: list = [Paragraph(escape(analysis.title), h1)]
    story.append(
        Paragraph(f"Generated {datetime.now(timezone.utc):%d %B %Y} by Starcode", small)
    )
    story.append(Spacer(1, 4 * mm))
    story.append(
        Table(
            [[Paragraph(f"<b>Important:</b> {escape(DISCLAIMER)}", small)]],
            colWidths=[document.width],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fffbeb")),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#f59e0b")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            ),
        )
    )
    story.append(Spacer(1, 5 * mm))

    confidence = report.confidence_summary or {}
    if confidence.get("overall") is not None:
        story.append(
            Paragraph(
                f"<b>Overall confidence:</b> {confidence['overall']:.0%} "
                f"({escape(str(confidence.get('band', '')))})",
                body,
            )
        )
        story.append(Paragraph(escape(str(confidence.get("rationale", ""))), small))
        story.append(Spacer(1, 4 * mm))

    story.append(Paragraph("Executive Summary", h2))
    story.append(Paragraph(escape(report.executive_summary), body))
    story.append(Spacer(1, 3 * mm))

    for section in report.sections:
        if section["key"] == "executive_summary":
            continue
        story.append(Paragraph(escape(section["title"]), h2))
        story.append(HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#cbd5e1")))

        if section.get("is_empty"):
            story.append(Paragraph(f"<i>{escape(section.get('empty_reason', ''))}</i>", small))
            continue

        for claim in section.get("claims", []):
            label = claim.get("presentation", {}).get("label", claim["claim_type"])
            colour = tone_colors.get(claim["claim_type"], "#475569")
            story.append(
                Paragraph(
                    f'<font color="{colour}"><b>[{escape(label)}]</b></font> '
                    f"{escape(claim['statement'])}",
                    body,
                )
            )
            if claim.get("quoted_text") and claim["claim_type"] == "source_text":
                story.append(Paragraph(escape(claim["quoted_text"][:1500]), quote))
            if claim.get("reasoning"):
                story.append(Paragraph(f"<i>{escape(claim['reasoning'])}</i>", small))
            basis = claim.get("confidence_basis") or ""
            story.append(
                Paragraph(
                    f"Confidence {claim['confidence']:.0%}"
                    + (f" — {escape(basis)}" if basis else ""),
                    small,
                )
            )
            for reference in claim.get("references", []):
                mark = "" if reference.get("verified") else " (unverified citation)"
                story.append(
                    Paragraph(f"• {escape(reference['citation_text'])}{escape(mark)}", small)
                )
            story.append(Spacer(1, 2 * mm))

        data = section.get("data") or {}
        if section["key"] == "timeline" and data.get("events"):
            rows = [["Date", "Event", "Certainty"]]
            for event in data["events"][:60]:
                rows.append(
                    [
                        Paragraph(escape(str(event.get("label", event.get("year")))), small),
                        Paragraph(escape(str(event.get("description", "")))[:300], small),
                        Paragraph(escape(str(event.get("certainty", ""))), small),
                    ]
                )
            table = Table(
                rows,
                colWidths=[document.width * 0.22, document.width * 0.58, document.width * 0.20],
                repeatRows=1,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 3 * mm))

        if section["key"] == "references" and data.get("references"):
            for reference in data["references"]:
                mark = "" if reference.get("verified") else " (unverified — check before citing)"
                story.append(
                    Paragraph(f"• {escape(reference['citation_text'])}{escape(mark)}", small)
                )
            story.append(Spacer(1, 3 * mm))

        if section["key"] == "further_reading" and data.get("items"):
            for item in data["items"]:
                story.append(
                    Paragraph(
                        f"• <b>{escape(item['title'])}</b> — {escape(item['citation'])}<br/>"
                        f"{escape(item['why'])}",
                        small,
                    )
                )
            story.append(Spacer(1, 3 * mm))

    story.append(PageBreak())
    story.append(Paragraph("How this report was produced", h2))
    story.append(
        Paragraph(
            "Each step below is a specialist agent. Deterministic steps use the platform's "
            "own calendar and ephemeris engines and consult no language model.",
            small,
        )
    )
    story.append(Spacer(1, 2 * mm))
    for step in report.reasoning_trace or []:
        story.append(
            Paragraph(
                f"<b>{step['step']}. {escape(step['agent'])}</b> "
                f"({escape(step['status'])}"
                + (f", {escape(str(step.get('model')))}" if step.get("model") else "")
                + f") — {escape(str(step.get('reasoning', '')))}",
                small,
            )
        )

    document.build(story)
    return buffer.getvalue()


# --------------------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------------------


def _build_docx(analysis: Analysis) -> bytes:
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # pragma: no cover
        raise ExportError("DOCX export requires python-docx, which is not installed.") from exc

    report = _report_of(analysis)
    document = docx.Document()

    document.core_properties.title = analysis.title
    document.core_properties.author = "Starcode"
    document.core_properties.comments = SHORT_DISCLAIMER

    document.add_heading(analysis.title, level=0)
    subtitle = document.add_paragraph(
        f"Generated {datetime.now(timezone.utc):%d %B %Y} by Starcode"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT

    warning = document.add_paragraph()
    run = warning.add_run(f"Important: {DISCLAIMER}")
    run.bold = True
    run.font.size = Pt(9)

    confidence = report.confidence_summary or {}
    if confidence.get("overall") is not None:
        paragraph = document.add_paragraph()
        paragraph.add_run(
            f"Overall confidence: {confidence['overall']:.0%} ({confidence.get('band', '')}). "
        ).bold = True
        paragraph.add_run(str(confidence.get("rationale", "")))

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(report.executive_summary)

    tone_colors = {
        "source_text": RGBColor(0x47, 0x55, 0x69),
        "verified_history": RGBColor(0x15, 0x80, 0x3D),
        "astronomical_calculation": RGBColor(0x1D, 0x4E, 0xD8),
        "textual_analysis": RGBColor(0x0F, 0x76, 0x6E),
        "traditional_interpretation": RGBColor(0xA1, 0x62, 0x07),
        "scholarly_interpretation": RGBColor(0x7C, 0x3A, 0xED),
        "ai_hypothesis": RGBColor(0xC2, 0x41, 0x0C),
        "uncertain": RGBColor(0x64, 0x74, 0x8B),
    }

    for section in report.sections:
        if section["key"] == "executive_summary":
            continue
        document.add_heading(section["title"], level=1)

        if section.get("is_empty"):
            document.add_paragraph(section.get("empty_reason", "")).italic = True
            continue

        for claim in section.get("claims", []):
            paragraph = document.add_paragraph()
            label = claim.get("presentation", {}).get("label", claim["claim_type"])
            tag = paragraph.add_run(f"[{label}] ")
            tag.bold = True
            tag.font.color.rgb = tone_colors.get(claim["claim_type"], RGBColor(0, 0, 0))
            paragraph.add_run(claim["statement"])

            if claim.get("quoted_text") and claim["claim_type"] == "source_text":
                quoted = document.add_paragraph(claim["quoted_text"][:1500])
                quoted.style = document.styles["Quote"] if "Quote" in [
                    s.name for s in document.styles
                ] else quoted.style

            if claim.get("reasoning"):
                note = document.add_paragraph()
                run = note.add_run(claim["reasoning"])
                run.italic = True
                run.font.size = Pt(9)

            meta = document.add_paragraph()
            run = meta.add_run(
                f"Confidence {claim['confidence']:.0%}"
                + (f" — {claim['confidence_basis']}" if claim.get("confidence_basis") else "")
            )
            run.font.size = Pt(8)

            for reference in claim.get("references", []):
                mark = "" if reference.get("verified") else " (unverified citation)"
                item = document.add_paragraph(
                    f"{reference['citation_text']}{mark}", style="List Bullet"
                )
                item.runs[0].font.size = Pt(8)

        data = section.get("data") or {}
        if section["key"] == "timeline" and data.get("events"):
            table = document.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            headers = table.rows[0].cells
            headers[0].text, headers[1].text, headers[2].text = "Date", "Event", "Certainty"
            for event in data["events"][:60]:
                cells = table.add_row().cells
                cells[0].text = str(event.get("label", event.get("year", "")))
                cells[1].text = str(event.get("description", ""))[:300]
                cells[2].text = str(event.get("certainty", ""))

        if section["key"] == "references" and data.get("references"):
            for reference in data["references"]:
                mark = "" if reference.get("verified") else " (unverified — check before citing)"
                document.add_paragraph(
                    f"{reference['citation_text']}{mark}", style="List Bullet"
                )

        if section["key"] == "further_reading" and data.get("items"):
            for item in data["items"]:
                document.add_paragraph(
                    f"{item['title']} — {item['citation']}. {item['why']}", style="List Bullet"
                )

    document.add_page_break()
    document.add_heading("How this report was produced", level=1)
    for step in report.reasoning_trace or []:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(f"{step['agent']} ({step['status']}): ").bold = True
        paragraph.add_run(str(step.get("reasoning", "")))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


__all__ = [
    "export_analysis",
    "EXPORT_FORMATS",
    "MIME_TYPES",
    "ExportError",
    "DISCLAIMER",
    "SHORT_DISCLAIMER",
]
